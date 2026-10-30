"""Version-bound approvals and immutable private DOCX exports."""
import hashlib,io
from typing import Literal
from .pdf_export import render_pdf
from docx import Document
from fastapi import HTTPException,Depends,Response
from pydantic import BaseModel,Field
from sqlalchemy import select,func
from .models import Application,ArtifactApproval,ArtifactExport,ResumeArtifact,Opportunity,User
from .career_gateway import digest,validate_claims


def lock_owner(db,owner):
    db.scalar(select(User).where(User.id==owner).with_for_update())


def application_for(db,opp):
    row=db.scalar(select(Application).where(Application.opportunity_id==opp.id))
    if row is None:
        row=Application(owner_id=opp.owner_id,opportunity_id=opp.id);db.add(row);db.flush()
    return row


def approval_reason(db,a,approval,current_hash):
    if approval is None:return 'Not approved'
    if approval.invalidated_reason:return approval.invalidated_reason
    latest=db.scalar(select(func.max(ResumeArtifact.version)).where(ResumeArtifact.opportunity_id==a.opportunity_id))
    if a.version!=latest:return 'A newer résumé version exists'
    if a.input_hash!=current_hash:return 'Profile or job requirements changed'
    if a.status!='approved' or approval.content_hash!=digest(a.claims) or a.content_hash!=approval.content_hash or a.approved_hash!=approval.content_hash:return 'Approved content changed'
    if validate_claims(a.claims,a.input_json['facts'],a.input_json.get('_claim_validation')):return 'Claim validation failed'
    return ''


def artifact_view(db,a,current_hash):
    approval=db.scalar(select(ArtifactApproval).where(ArtifactApproval.artifact_id==a.id))
    reason=approval_reason(db,a,approval,current_hash)
    return dict(id=a.id,version=a.version,status='needs_review' if a.input_hash!=current_hash else a.status,content_hash=a.content_hash,approved_hash=a.approved_hash,claims=a.claims,evidence=a.input_json['facts'],findings=a.findings,stale=a.input_hash!=current_hash,approval_valid=not reason,approval_id=approval.id if approval else None,generation=a.input_json.get('_generation_audit',{}),verification=a.input_json.get('_claim_validation',{}))


def invalidate(db,opp,reason):
    app=db.scalar(select(Application).where(Application.opportunity_id==opp.id))
    if app and app.current_approval_id:
        approval=db.get(ArtifactApproval,app.current_approval_id)
        if approval:approval.invalidated_reason=reason
        app.status='needs_review'


def save_approval(db,opp,a,owner):
    app=application_for(db,opp)
    approval=db.scalar(select(ArtifactApproval).where(ArtifactApproval.artifact_id==a.id))
    if approval is None:
        approval=ArtifactApproval(application_id=app.id,artifact_id=a.id,actor_id=owner,content_hash=a.content_hash,input_hash=a.input_hash);db.add(approval);db.flush()
    if approval.invalidated_reason:raise HTTPException(409,'This approval was invalidated; create a new résumé version')
    app.current_approval_id=approval.id;app.status='ready'
    return app,approval


def render_docx(claims):
    # Only approved claim text enters the exported body. Labels and source
    # excerpts remain in the application, not in the candidate's résumé.
    doc=Document();doc.core_properties.title='Résumé';doc.core_properties.author=''
    for claim in claims:doc.add_paragraph(claim['text'])
    out=io.BytesIO();doc.save(out);data=out.getvalue()
    if [p.text for p in Document(io.BytesIO(data)).paragraphs]!=[c['text'] for c in claims]:raise ValueError('Rendered résumé does not match approved content')
    return data

class ExportRequest(BaseModel):
    format:Literal['docx','pdf']='docx'
    version:int=Field(ge=1)
    content_hash:str=Field(min_length=64,max_length=64)


def install(app,current_user,db_session,snapshot,get_opp):
    def owned_artifact(db,aid,owner):
        a=db.scalar(select(ResumeArtifact).join(Opportunity).where(ResumeArtifact.id==aid,Opportunity.owner_id==owner))
        if not a:raise HTTPException(404,'Artifact not found')
        return a,get_opp(db,a.opportunity_id,owner)

    def application_summary(db,row):
        opp=db.get(Opportunity,row.opportunity_id);approval=db.get(ArtifactApproval,row.current_approval_id) if row.current_approval_id else None
        a=db.get(ResumeArtifact,approval.artifact_id) if approval else None
        reason=approval_reason(db,a,approval,digest(snapshot(db,opp))) if a else 'Not approved'
        from .tracking_models import ApplicationTracking
        tracking=db.get(ApplicationTracking,row.id)
        artifact_status='needs_review' if approval and reason else row.status
        return dict(id=row.id,opportunity_id=opp.id,title=opp.title,status=tracking.status if tracking else artifact_status,artifact_status=artifact_status,approved_artifact_id=a.id if a else None,approved_version=a.version if a else None,approval_valid=not reason)

    @app.get('/applications')
    def applications(owner=Depends(current_user),db=Depends(db_session)):
        return [application_summary(db,a) for a in db.scalars(select(Application).where(Application.owner_id==owner).order_by(Application.created_at.desc()))]

    @app.get('/applications/{appid}')
    def application_detail(appid:str,owner=Depends(current_user),db=Depends(db_session)):
        row=db.scalar(select(Application).where(Application.id==appid,Application.owner_id==owner))
        if not row:raise HTTPException(404,'Application not found')
        opp=get_opp(db,row.opportunity_id,owner);key=digest(snapshot(db,opp));result=application_summary(db,row)
        approval=db.get(ArtifactApproval,row.current_approval_id) if row.current_approval_id else None
        a=db.get(ResumeArtifact,approval.artifact_id) if approval else None
        result.update(job_snapshot={'title':opp.title,'description':opp.description},approval=None,approved_artifact=artifact_view(db,a,key) if a else None,artifact_versions=[artifact_view(db,v,key) for v in db.scalars(select(ResumeArtifact).where(ResumeArtifact.opportunity_id==opp.id).order_by(ResumeArtifact.version.desc()))],exports=[])
        if approval:
            reason=approval_reason(db,a,approval,key)
            result['approval']=dict(id=approval.id,artifact_id=a.id,version=a.version,content_hash=approval.content_hash,approved_at=approval.approved_at,valid=not reason,invalidated_reason=reason)
            result['exports']=[dict(id=e.id,format=e.format,content_hash=e.content_hash,file_hash=e.file_hash,created_at=e.created_at) for e in db.scalars(select(ArtifactExport).where(ArtifactExport.approval_id==approval.id))]
        return result

    @app.get('/artifacts/{aid}')
    def artifact_detail(aid:str,owner=Depends(current_user),db=Depends(db_session)):
        a,opp=owned_artifact(db,aid,owner);return artifact_view(db,a,digest(snapshot(db,opp)))

    @app.post('/artifacts/{aid}/export')
    def export(aid:str,body:ExportRequest,owner=Depends(current_user),db=Depends(db_session)):
        lock_owner(db,owner);a,opp=owned_artifact(db,aid,owner)
        approval=db.scalar(select(ArtifactApproval).where(ArtifactApproval.artifact_id==aid))
        reason=approval_reason(db,a,approval,digest(snapshot(db,opp)))
        if reason:raise HTTPException(409,reason)
        if a.version!=body.version or a.content_hash!=body.content_hash:raise HTTPException(409,'Requested version does not match approved content')
        row=db.scalar(select(ArtifactExport).where(ArtifactExport.approval_id==approval.id,ArtifactExport.format==body.format))
        if row is None:
            try:data=render_pdf(a.claims) if body.format=='pdf' else render_docx(a.claims)
            except Exception:raise HTTPException(503,'Export rendering failed; retry safely')
            row=ArtifactExport(approval_id=approval.id,format=body.format,renderer_version=body.format+'-v1',content_hash=approval.content_hash,file_hash=hashlib.sha256(data).hexdigest(),file_bytes=data);db.add(row);db.flush()
        if hashlib.sha256(row.file_bytes).hexdigest()!=row.file_hash or row.content_hash!=approval.content_hash:raise HTTPException(409,'Export integrity check failed')
        data=row.file_bytes;file_hash=row.file_hash;db.commit()
        return Response(data,media_type='application/pdf' if body.format=='pdf' else 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',headers={'Content-Disposition':f'attachment; filename="resume-v{a.version}.{body.format}"','Cache-Control':'private, no-store','X-Artifact-Content-Hash':approval.content_hash,'X-File-SHA256':file_hash})
