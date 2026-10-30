"""Local source-bound adapter. No LLM, inferred claims, or independent verification."""
import io, re, zipfile
from pypdf import PdfReader
from docx import Document
MAX_EXPANDED = 30 * 1024 * 1024

def inspect_file(data: bytes, filename: str):
    if filename.lower().endswith('.pdf') and data.startswith(b'%PDF-'): return
    if filename.lower().endswith('.docx') and data.startswith(b'PK'):
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if sum(i.file_size for i in z.infolist()) > MAX_EXPANDED: raise ValueError('Document expands beyond the limit')
            if 'word/document.xml' not in z.namelist(): raise ValueError('Not a DOCX document')
        return
    raise ValueError('Upload a text-based PDF or DOCX')

def parse(data: bytes, filename: str):
    inspect_file(data, filename)
    if filename.lower().endswith('.pdf'):
        reader = PdfReader(io.BytesIO(data))
        if len(reader.pages) > 40: raise ValueError('Maximum 40 pages')
        text = '\n'.join(p.extract_text() or '' for p in reader.pages)
    else:
        doc = Document(io.BytesIO(data))
        text = '\n'.join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells])
    if not text.strip(): raise ValueError('No readable text; scanned documents need OCR')
    if len(text) > 200_000: raise ValueError('Extracted text is too large')
    return text

def extract(text):
    for m in re.finditer(r'[^\n]+', text):
        value = m.group().strip()
        if len(value) < 12: continue
        category = 'email' if re.fullmatch(r'[^\s@]+@[^\s@]+\.[^\s@]+', value) else 'statement'
        offset = m.start() + len(m.group()) - len(m.group().lstrip())
        yield dict(value=value, category=category, source=value, start=offset, end=offset+len(value))
