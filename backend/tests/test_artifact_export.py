import io,hashlib
import pytest
from docx import Document
from sqlalchemy import select,func
from test_profile import env,client
from test_career import confirmed,opportunity,draft
from app.models import Application,ArtifactApproval,ArtifactExport
from app import artifact_export

def setup(c):
    confirmed(c);oid=opportunity(c);a=draft(c,oid)
    approval=c.post(f"/artifacts/{a['id']}/approve",json={'version':a['version']})
    assert approval.status_code==200
    return oid,a,approval.json()

def export(c,a):return c.post(f"/artifacts/{a['id']}/export",json={'version':a['version'],'content_hash':a['content_hash']})

def test_application_approved_version_export_and_provenance(env):
    c=client();oid,a,approval=setup(c)
    record=c.get('/applications/'+approval['application_id']).json()
    assert record['approval']['valid'] and record['approved_version']==1
    assert record['approved_artifact']['claims']==a['claims']
    assert record['approved_artifact']['evidence']==a['evidence']
    result=export(c,a);assert result.status_code==200
    assert [p.text for p in Document(io.BytesIO(result.content)).paragraphs]==[x['text'] for x in a['claims']]
    assert result.headers['x-artifact-content-hash']==a['content_hash']
    assert result.headers['x-file-sha256']==hashlib.sha256(result.content).hexdigest()
    assert export(c,a).content==result.content
    assert c.post(f"/artifacts/{a['id']}/approve",json={'version':1}).json()['approval_id']==approval['approval_id']
    with env() as db:
        assert db.scalar(select(func.count()).select_from(ArtifactApproval))==1
        assert db.scalar(select(func.count()).select_from(ArtifactExport))==1

def test_edit_invalidates_approval_preserves_history(env):
    c=client();oid,a,approval=setup(c);old_bytes=export(c,a).content
    claims=list(reversed(a['claims']))
    new=c.post(f"/artifacts/{a['id']}/edit",json={'version':1,'claims':claims}).json()
    assert export(c,a).status_code==409
    record=c.get('/applications/'+approval['application_id']).json()
    assert not record['approval']['valid'] and 'edited' in record['approval']['invalidated_reason']
    historical=c.get('/artifacts/'+a['id']).json();assert historical['claims']==a['claims'] and historical['evidence']==a['evidence']
    b=c.get('/artifacts/'+new['id']).json();assert export(c,b).status_code==409
    assert c.post('/artifacts/'+b['id']+'/approve',json={'version':2}).status_code==200
    output=export(c,b);assert output.status_code==200
    assert [p.text for p in Document(io.BytesIO(output.content)).paragraphs]==[x['text'] for x in claims]
    with env() as db:
        old=db.scalar(select(ArtifactExport).join(ArtifactApproval).where(ArtifactApproval.artifact_id==a['id']))
        assert old.file_bytes==old_bytes

def test_export_rejects_unapproved_stale_wrong_version_and_foreign(env):
    c=client();oid,a,approval=setup(c)
    b=client('foreign@example.test')
    assert b.get('/applications').json()==[]
    for path in ['/applications/'+approval['application_id'],'/artifacts/'+a['id']]:assert b.get(path).status_code==404
    assert export(b,a).status_code==404
    assert c.post('/artifacts/'+a['id']+'/export',json={'version':2,'content_hash':a['content_hash']}).status_code==409
    assert c.post('/artifacts/'+a['id']+'/export',json={'version':1,'content_hash':'0'*64}).status_code==409
    p=c.get('/proposals').json()[0];c.post('/proposals/'+p['id']+'/review',json={'revision':p['revision'],'action':'rejected'})
    assert export(c,a).status_code==409
    assert not c.get('/applications/'+approval['application_id']).json()['approval_valid']

def test_render_failure_is_atomic_and_recoverable(env,monkeypatch):
    c=client();oid,a,approval=setup(c);render=artifact_export.render_docx
    monkeypatch.setattr(artifact_export,'render_docx',lambda _: (_ for _ in ()).throw(RuntimeError('synthetic render failure')))
    assert export(c,a).status_code==503
    with env() as db:assert db.scalar(select(func.count()).select_from(ArtifactExport))==0
    monkeypatch.setattr(artifact_export,'render_docx',render);assert export(c,a).status_code==200

def test_corrupt_stored_export_fails_closed(env):
    c=client();oid,a,approval=setup(c);export(c,a)
    with env.begin() as db:db.scalar(select(ArtifactExport)).file_bytes=b'corrupt'
    assert export(c,a).status_code==409

def test_concurrent_exports_are_idempotent_postgres(env):
    import os,threading
    if not os.getenv('TEST_DATABASE_URL'):pytest.skip('PostgreSQL concurrency')
    c=client();oid,a,approval=setup(c);results=[]
    threads=[threading.Thread(target=lambda:results.append(export(c,a))) for _ in range(2)]
    for t in threads:t.start()
    for t in threads:t.join(10)
    assert len(results)==2 and all(r.status_code==200 for r in results)
    assert results[0].content==results[1].content
    with env() as db:assert db.scalar(select(func.count()).select_from(ArtifactExport))==1

def test_export_then_edit_serializes_and_only_exports_approved_text(env,monkeypatch):
    import os,threading,time
    if not os.getenv('TEST_DATABASE_URL'):pytest.skip('PostgreSQL race ordering')
    c=client();oid,a,approval=setup(c)
    entered=threading.Event();release=threading.Event();edit_done=threading.Event();results={}
    original=artifact_export.render_docx
    def blocked_render(claims):
        entered.set();assert release.wait(10);return original(claims)
    monkeypatch.setattr(artifact_export,'render_docx',blocked_render)
    claims=[{'text':'Unsupported new text must not be exported','fact_ids':a['claims'][0]['fact_ids']}]
    exporter=threading.Thread(target=lambda:results.update(export=export(c,a)))
    def edit():
        results['edit']=c.post('/artifacts/'+a['id']+'/edit',json={'version':1,'claims':claims});edit_done.set()
    editor=threading.Thread(target=edit)
    exporter.start();assert entered.wait(5);editor.start();time.sleep(.15)
    assert not edit_done.is_set()
    release.set();exporter.join(10);editor.join(10)
    assert results['export'].status_code==200 and results['edit'].status_code==200
    assert [p.text for p in Document(io.BytesIO(results['export'].content)).paragraphs]==[x['text'] for x in a['claims']]
    assert export(c,a).status_code==409
    updated=c.get('/artifacts/'+results['edit'].json()['id']).json();assert export(c,updated).status_code==409
    assert not c.get('/applications/'+approval['application_id']).json()['approval_valid']

def test_edit_then_export_rejects_invalidated_approval(env,monkeypatch):
    import os,threading,time
    from app import career
    if not os.getenv('TEST_DATABASE_URL'):pytest.skip('PostgreSQL race ordering')
    c=client();oid,a,approval=setup(c)
    entered=threading.Event();release=threading.Event();export_done=threading.Event();results={}
    original=career.validate_claims
    def blocked_validation(*args):
        entered.set();assert release.wait(10);return original(*args)
    monkeypatch.setattr(career,'validate_claims',blocked_validation)
    editor=threading.Thread(target=lambda:results.update(edit=c.post('/artifacts/'+a['id']+'/edit',json={'version':1,'claims':list(reversed(a['claims']))})))
    def do_export():results['export']=export(c,a);export_done.set()
    exporter=threading.Thread(target=do_export)
    editor.start();assert entered.wait(5);exporter.start();time.sleep(.15);assert not export_done.is_set()
    release.set();editor.join(10);exporter.join(10)
    assert results['edit'].status_code==200 and results['export'].status_code==409
    with env() as db:assert db.scalar(select(func.count()).select_from(ArtifactExport))==0

def test_approve_then_edit_invalidates_saved_approval_atomically(env,monkeypatch):
    import os,threading,time
    from app import career
    if not os.getenv('TEST_DATABASE_URL'):pytest.skip('PostgreSQL race ordering')
    c=client();confirmed(c);oid=opportunity(c);a=draft(c,oid)
    entered=threading.Event();release=threading.Event();edit_done=threading.Event();results={}
    original=career.save_approval
    def blocked_save(*args):
        entered.set();assert release.wait(10);return original(*args)
    monkeypatch.setattr(career,'save_approval',blocked_save)
    approver=threading.Thread(target=lambda:results.update(approve=c.post('/artifacts/'+a['id']+'/approve',json={'version':1})))
    def edit():results['edit']=c.post('/artifacts/'+a['id']+'/edit',json={'version':1,'claims':list(reversed(a['claims']))});edit_done.set()
    editor=threading.Thread(target=edit)
    approver.start();assert entered.wait(5);editor.start();time.sleep(.15);assert not edit_done.is_set()
    release.set();approver.join(10);editor.join(10)
    assert results['approve'].status_code==200 and results['edit'].status_code==200
    record=c.get('/applications/'+results['approve'].json()['application_id']).json()
    assert not record['approval_valid'] and 'edited' in record['approval']['invalidated_reason']
    assert export(c,a).status_code==409
