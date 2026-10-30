import io, time
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select, func
from docx import Document as WordDocument
from app import db as database
from app import main, worker
from app.models import Base, Document, Job, Proposal, FactVersion, ReviewEvent

@pytest.fixture()
def env(tmp_path, monkeypatch):
    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker
    import os
    url = os.getenv('TEST_DATABASE_URL', f'sqlite:///{tmp_path}/test.db')
    if os.getenv('TEST_DATABASE_URL'):
        from sqlalchemy.engine import make_url
        assert (make_url(url).database or '').endswith('_test'), 'Use a disposable database ending in _test'
    engine = create_engine(url, connect_args={'check_same_thread':False} if url.startswith('sqlite') else {})
    Base.metadata.drop_all(engine); Base.metadata.create_all(engine)
    session = sessionmaker(engine, expire_on_commit=False)
    monkeypatch.setattr(main,'Session',session); monkeypatch.setattr(worker,'Session',session)
    monkeypatch.setenv('EXTRACTION_ADAPTER','deterministic')
    monkeypatch.setenv('STORAGE_ROOT',str(tmp_path/'storage'))
    yield session
    engine.dispose()

def client(email='alex@example.test'):
    c=TestClient(main.app, headers={'Origin':'http://localhost:3000'})
    assert c.post('/auth/signup',json={'email':email,'password':'synthetic-passphrase-2026'}).status_code==200
    return c

def resume(email='alex@example.test'):
    doc=WordDocument()
    for line in ['Alex Morgan',email,'Senior data analyst at Cedar Labs, 2021–2025','Reduced reporting time by 30% through Python automation.','BSc Statistics, North Coast University, 2020']:
        doc.add_paragraph(line)
    out=io.BytesIO();doc.save(out);return out.getvalue()

def upload(c,data=None):
    return c.post('/documents',files={'file':('resume.docx',data or resume(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})

def test_full_review_and_provenance(env):
    c=client(); assert upload(c).status_code==201; assert worker.run_once()
    p=c.get('/proposals').json(); assert len(p)>=4
    for i,action in enumerate(['accepted','corrected','rejected','deferred']):
        r=c.post(f"/proposals/{p[i]['id']}/review",json={'action':action,'revision':0,'value':'Candidate corrected statement'})
        assert r.status_code==200
    facts=c.get('/profile').json(); assert len(facts)==2
    assert {f['provenance'] for f in facts}=={'candidate_correction','source_supported'}
    assert all(f['independently_verified'] is False for f in facts)
    revised=c.get('/proposals').json()[1]; assert revised['source']==p[1]['source']
    assert c.post(f"/proposals/{p[0]['id']}/review",json={'action':'rejected','revision':0}).status_code==409
    assert c.post(f"/proposals/{p[0]['id']}/review",json={'action':'corrected','value':'Updated candidate statement','revision':1}).status_code==200
    with env() as db:
        assert db.scalar(select(func.count()).select_from(FactVersion))==3
        assert db.scalar(select(func.count()).select_from(ReviewEvent))==5

def test_isolation_and_session(env):
    a=client();upload(a);worker.run_once();b=client('other@example.test')
    assert b.get('/documents').json()==[];assert b.get('/proposals').json()==[];assert b.get('/profile').json()==[]
    p=a.get('/proposals').json()[0];d=a.get('/documents').json()[0]
    assert b.post(f"/proposals/{p['id']}/review",json={'action':'accepted','revision':0}).status_code==404
    assert b.post(f"/documents/{d['id']}/retry").status_code==404
    assert upload(b).json()['id']!=d['id']
    assert a.post('/logout').status_code==200;assert a.get('/me').status_code==401
    assert a.post('/auth/signin',json={'email':'alex@example.test','password':'wrong-password-long'}).status_code==401
    assert a.post('/auth/signin',json={'email':'alex@example.test','password':'synthetic-passphrase-2026'}).status_code==200
    assert a.post('/logout',headers={'Origin':'https://evil.example'}).status_code==403

def test_invalid_uploads_and_duplicates(env):
    c=client();assert upload(c,b'not a word document').status_code==422
    assert upload(c,b'x'*(5*1024*1024+1)).status_code==413
    assert c.get('/documents').json()==[]
    content=resume();first=upload(c,content).json();second=upload(c,content).json()
    assert first['id']==second['id'];assert second['duplicate']
    assert worker.run_once();assert not worker.run_once()
    with env.begin() as db:
        job=db.scalar(select(Job));job.status='queued'
        count=db.scalar(select(func.count()).select_from(Proposal))
    assert worker.run_once()
    with env() as db:assert db.scalar(select(func.count()).select_from(Proposal))==count

def test_failed_job_and_retry(env,monkeypatch):
    c=client();upload(c);original=worker.parse
    monkeypatch.setattr(worker,'parse',lambda *args: (_ for _ in ()).throw(ValueError('Synthetic parser failure')))
    worker.run_once();d=c.get('/documents').json()[0];assert d['status']=='failed';assert d['error']=='Synthetic parser failure'
    assert c.post(f"/documents/{d['id']}/retry").status_code==200
    monkeypatch.setattr(worker,'parse',original);worker.run_once();assert c.get('/documents').json()[0]['status']=='ready'

def test_crashed_worker_lease_recovery(env):
    c=client();upload(c)
    with env.begin() as db:
        j=db.scalar(select(Job));j.status='running';j.lease_until=time.time()+30;j.lease_token='dead-worker'
    assert not worker.run_once()
    with env.begin() as db: db.scalar(select(Job)).lease_until=time.time()-1
    assert worker.run_once();assert c.get('/documents').json()[0]['status']=='ready'

def test_conflicting_email_passages(env):
    c=client();upload(c);worker.run_once();upload(c,resume('alex.new@example.test'));worker.run_once()
    conflicts=[p for p in c.get('/proposals').json() if p['conflicts']]
    assert len(conflicts)==2;assert c.get('/profile').json()==[]

def test_transient_failures_exhaust_and_recover(env,monkeypatch):
    c=client();upload(c);original=worker.parse
    monkeypatch.setattr(worker,'parse',lambda *args: (_ for _ in ()).throw(RuntimeError('private failure detail')))
    for _ in range(3):
        with env.begin() as db: db.scalar(select(Job)).available_at=0
        worker.run_once()
    d=c.get('/documents').json()[0];assert d['status']=='failed';assert d['attempts']==3
    assert 'private failure' not in d['error']
    monkeypatch.setattr(worker,'parse',original);c.post(f"/documents/{d['id']}/retry");worker.run_once()
    assert c.get('/documents').json()[0]['status']=='ready'

def test_stale_worker_cannot_commit(env,monkeypatch):
    c=client();upload(c);original=worker.parse
    def lose_lease(*args):
        with env.begin() as db:db.scalar(select(Job)).lease_token='replacement-worker'
        return original(*args)
    monkeypatch.setattr(worker,'parse',lose_lease);worker.run_once()
    assert c.get('/proposals').json()==[]

def test_empty_docx(env):
    c=client();doc=WordDocument();buf=io.BytesIO();doc.save(buf)
    upload(c,buf.getvalue());worker.run_once();assert c.get('/documents').json()[0]['status']=='failed'

def test_storage_failure_creates_no_job(env,monkeypatch):
    from pathlib import Path
    c=client();original=Path.open
    def fail_write(self,*args,**kwargs):
        if args and args[0]=='xb':raise OSError('disk unavailable')
        return original(self,*args,**kwargs)
    monkeypatch.setattr(Path,'open',fail_write)
    assert upload(c).status_code==503
    assert c.get('/documents').json()==[]
    with env() as db:assert db.scalar(select(func.count()).select_from(Job))==0

def test_similar_metric_conflict(env):
    c=client();upload(c);worker.run_once()
    doc=WordDocument();doc.add_paragraph('Reduced reporting time by 25% through Python automation.')
    buf=io.BytesIO();doc.save(buf);upload(c,buf.getvalue());worker.run_once()
    conflicts=[p for p in c.get('/proposals').json() if p['conflicts']]
    assert len(conflicts)==2

def test_concurrent_workers_on_postgres(env):
    from concurrent.futures import ThreadPoolExecutor
    if env.kw['bind'].dialect.name!='postgresql':pytest.skip('requires PostgreSQL row locks')
    c=client();upload(c)
    with ThreadPoolExecutor(max_workers=2) as pool:results=list(pool.map(lambda _:worker.run_once(),range(2)))
    assert results.count(True)==1
    assert c.get('/documents').json()[0]['status']=='ready'
    with env() as db:assert db.scalar(select(Job)).attempts==1

def test_repeated_worker_crashes_become_retryable_failure(env):
    c=client();upload(c)
    with env.begin() as db:
        j=db.scalar(select(Job));j.status='running';j.lease_until=0;j.attempts=3
    worker.run_once();d=c.get('/documents').json()[0];assert d['status']=='failed'
    assert c.post(f"/documents/{d['id']}/retry").status_code==200
    worker.run_once();assert c.get('/documents').json()[0]['status']=='ready'

def test_pdf_source_offsets(env):
    c=client()
    content=b'BT /F1 12 Tf 50 700 Td (Built Python workflows for Cedar Labs.) Tj ET'
    objs=[b'<< /Type /Catalog /Pages 2 0 R >>',b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>',b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',b'<< /Length '+str(len(content)).encode()+b' >>\nstream\n'+content+b'\nendstream']
    data=b'%PDF-1.4\n';offsets=[0]
    for i,obj in enumerate(objs,1):
        offsets.append(len(data));data+=f'{i} 0 obj\n'.encode()+obj+b'\nendobj\n'
    xref=len(data);data+=b'xref\n0 6\n0000000000 65535 f \n'
    for offset in offsets[1:]:data+=f'{offset:010d} 00000 n \n'.encode()
    data+=f'trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF'.encode()
    assert c.post('/documents',files={'file':('resume.pdf',data,'application/pdf')}).status_code==201
    worker.run_once();p=c.get('/proposals').json()[0]
    with env() as db:
        text=db.scalar(select(Document)).text
        assert text[p['start']:p['end']]==p['source']==p['value']

def test_ai_gateway_worker_atomic_failure_retry_and_audit(env,monkeypatch):
    import json
    import httpx2
    from openai import OpenAI
    from app.ai_gateway import extract_facts, passages_for
    from app.models import GenerationRun
    c=client();upload(c)
    invalid=True
    def respond(request):
        body=json.loads(request.content)
        span=json.loads(body['input'])['untrusted_resume_passages'][0]
        fact={'category':'statement','span_id':'invented' if invalid else span['span_id'],'quote':span['text'],'value':span['text'],'fields':[]}
        return httpx2.Response(200,json={'id':'test','object':'response','created_at':1,'model':'test-model','status':'completed','output':[{'id':'message','type':'message','role':'assistant','status':'completed','content':[{'type':'output_text','text':json.dumps({'facts':[fact]}),'annotations':[]}]}],'usage':{'input_tokens':10,'output_tokens':10,'total_tokens':20}})
    with OpenAI(api_key='test-only',http_client=httpx2.Client(transport=httpx2.MockTransport(respond)),max_retries=0) as sdk:
        monkeypatch.setattr(worker,'extract_facts',lambda text:extract_facts(text,client=sdk,adapter='openai',model='test-model'))
        worker.run_once();assert c.get('/proposals').json()==[]
        d=c.get('/documents').json()[0];assert d['status']=='failed'
        invalid=False;c.post(f"/documents/{d['id']}/retry");worker.run_once()
    p=c.get('/proposals').json();assert len(p)==1 and p[0]['extraction_adapter']=='openai'
    with env() as db:
        runs=db.scalars(select(GenerationRun).order_by(GenerationRun.created_at)).all()
        assert [r.status for r in runs]==['failed','completed']
        assert runs[0].error_code=='invalid_citation'
        assert runs[1].metadata_json['model']=='test-model'
        assert all('alex@example.test' not in json.dumps(r.metadata_json) for r in runs)
