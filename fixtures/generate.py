"""Generate synthetic documents; all organizations and achievements are fictional."""
from pathlib import Path
from docx import Document
root=Path(__file__).parent
for filename,email,achievement in [
 ('alex-morgan.docx','alex.morgan@example.test','Reduced monthly reporting time by 30% by automating SQL and Python workflows.'),
 ('alex-morgan-conflict.docx','alex.changed@example.test','Reduced monthly reporting time by 25% by automating SQL and Python workflows.')]:
    d=Document();d.add_heading('Alex Morgan',0)
    for line in [email,'Portland, Oregon','Senior Data Analyst | Cedar Labs | March 2021–August 2025',achievement,'Partnered with six operations teams to define quality metrics.','Data Analyst | Harbor Metrics | June 2018–February 2021','Built dashboard validation checks using Python and PostgreSQL.','BSc Statistics | North Coast University | 2018','Skills: Python, SQL, PostgreSQL, data visualization']:
        d.add_paragraph(line)
    d.save(root/filename)
(root/'invalid-resume.pdf').write_text('This deliberately invalid fixture is not a PDF.')
